"""Document discovery and raw text extraction."""

import logging
from dataclasses import dataclass
from pathlib import Path
from typing import List, Tuple

from docx import Document
from PyPDF2 import PdfReader


LOGGER = logging.getLogger(__name__)
SUPPORTED_EXTENSIONS = {".txt", ".pdf", ".docx"}


class DocumentReadError(Exception):
    """Raised when a document cannot be read or is not supported."""


@dataclass(frozen=True)
class ReadDocument:
    """Raw text and source metadata for one successfully read document."""

    source_path: Path
    text: str


def _read_text_file(path: Path) -> str:
    return path.read_text(encoding="utf-8")


def _read_pdf_file(path: Path) -> str:
    reader = PdfReader(str(path))
    return "\n".join(page.extract_text() or "" for page in reader.pages)


def _read_docx_file(path: Path) -> str:
    document = Document(str(path))
    return "\n".join(paragraph.text for paragraph in document.paragraphs)


def read_document(path: Path) -> ReadDocument:
    """Extract text from one supported document."""
    if not path.is_file():
        raise DocumentReadError(f"Document does not exist: {path}")

    extension = path.suffix.lower()
    readers = {
        ".txt": _read_text_file,
        ".pdf": _read_pdf_file,
        ".docx": _read_docx_file,
    }
    reader = readers.get(extension)
    if reader is None:
        raise DocumentReadError(f"Unsupported document type: {extension or '<none>'}")

    try:
        text = reader(path).strip()
    except Exception as exc:
        raise DocumentReadError(f"Could not read {path.name}: {exc}") from exc

    return ReadDocument(source_path=path, text=text)


def read_documents(data_dir: Path) -> Tuple[List[ReadDocument], List[str]]:
    """Read all supported files, returning successes and human-readable errors."""
    if not data_dir.is_dir():
        LOGGER.warning("Data directory does not exist: %s", data_dir)
        return [], [f"Data directory does not exist: {data_dir}"]

    documents: List[ReadDocument] = []
    errors: List[str] = []
    for path in sorted(data_dir.iterdir()):
        if not path.is_file() or path.suffix.lower() not in SUPPORTED_EXTENSIONS:
            continue
        try:
            documents.append(read_document(path))
        except DocumentReadError as exc:
            LOGGER.error("Skipping %s: %s", path.name, exc)
            errors.append(str(exc))
    return documents, errors